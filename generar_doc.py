import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Margenes
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def add_header_callout(text, bg_hex="6B21A8", text_color=RGBColor(255, 255, 255)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(107, 33, 168)

# Titulo principal
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(2)
run_title = title_p.add_run("PROYECTO NEXUS")
run_title.bold = True
run_title.font.size = Pt(24)
run_title.font.color.rgb = RGBColor(107, 33, 168)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(20)
run_sub = sub_p.add_run("Guía de Exposición y Documentación Técnica de Arquitectura")
run_sub.font.size = Pt(12)
run_sub.font.color.rgb = RGBColor(100, 100, 100)

intro_p = doc.add_paragraph()
intro_p.paragraph_format.space_after = Pt(12)
intro_p.add_run("Este documento contiene el desglose detallado para la exposición del proyecto NEXUS. Cada sección contiene el contenido sugerido para las diapositivas, el guion explicativo recomendado y la ubicación exacta de los componentes en el código fuente.")

secciones = [
    {
        "num": "1",
        "tema": "Autenticación con Firebase",
        "slides": [
            "Motor: Firebase Authentication (firebase/auth).",
            "Mecanismos: Registro (createUserWithEmailAndPassword), Login (signInWithEmailAndPassword), Logout (signOut).",
            "Persistencia Local: Integración con @react-native-async-storage/async-storage en React Native para no perder la sesión tras reiniciar la app.",
            "Estado Reactivo Global: AuthContext centraliza el listener onAuthStateChanged que conmuta automáticamente entre pantallas públicas y privadas."
        ],
        "speech": "Para la seguridad y gestión de identidades implementamos Firebase Authentication. La inicialización se realizó en firebaseConfig.js, donde configuramos la persistencia mediante AsyncStorage en entornos móviles para mantener activa la sesión aunque el usuario cierre la aplicación.\n\nTodo el estado de autenticación se gestiona de forma global y reactiva a través de AuthContext.js con el listener onAuthStateChanged. Esto significa que cuando el usuario inicia sesión en LoginScreen o crea su cuenta en RegisterScreen, no necesitamos forzar manualmente la navegación; App.js detecta automáticamente si existe un objeto user activo y conmuta de inmediato entre las pantallas públicas y el menú principal de la aplicación.",
        "codigo": [
            ("Configuración y persistencia", "firebaseConfig.js (Líneas 20-36)"),
            ("Manejador de estado global", "context/AuthContext.js (Líneas 14-26)"),
            ("Inicio de sesión", "screens/LoginScreen.js (Líneas 20-36)"),
            ("Registro de nuevo usuario", "screens/RegisterScreen.js (Líneas 44-52)"),
            ("Navegación protegida reactiva", "App.js (Líneas 109-126)")
        ]
    },
    {
        "num": "2",
        "tema": "Conexión y Persistencia con Cloud Firestore",
        "slides": [
            "Base de datos: Cloud Firestore (NoSQL orientada a colecciones y documentos).",
            "Colección 'usuarios': Guarda nombre, nombre preferido, email y fecha de creación.",
            "Colección 'recordatorios': Alarmas y recordatorios asociados al UID del usuario.",
            "Sincronización en tiempo real: Uso de onSnapshot para actualizar la UI instantáneamente.",
            "Operaciones CRUD: addDoc, setDoc, updateDoc, deleteDoc y filtros con where('uid', '==', user.uid)."
        ],
        "speech": "La persistencia de los datos del asistente se realiza con Cloud Firestore. Al registrar un usuario en RegisterScreen, se crea un documento en la colección 'usuarios' usando su identificador único user.uid como clave primaria.\n\nPara la gestión de recordatorios y alarmas creamos el RemindersContext.js. Aquí se ejecuta una consulta con query() filtrando únicamente los datos del usuario actual con where('uid', '==', user.uid). En lugar de hacer peticiones estáticas periódicas, usamos onSnapshot, lo que permite que cualquier creación con addDoc, actualización con updateDoc o eliminación con deleteDoc se sincronice en tiempo real con la interfaz gráfica de HomeScreen y RemindersScreen.",
        "codigo": [
            ("Instancia de base de datos", "firebaseConfig.js (Línea 37)"),
            ("Almacenamiento de perfil", "screens/RegisterScreen.js (Líneas 53-60)"),
            ("Lectura en tiempo real del perfil", "screens/HomeScreen.js (Líneas 30-41)"),
            ("Suscripción reactiva a recordatorios", "context/RemindersContext.js (Líneas 74-98)"),
            ("Funciones de creación, edición y borrado", "context/RemindersContext.js (Líneas 135-191)")
        ]
    },
    {
        "num": "3",
        "tema": "Implementación de Notificaciones Locales",
        "slides": [
            "Librería: expo-notifications (con arquitectura modular).",
            "Canales Android (Channels): 'nexus-recordatorios' (prioridad alta) y 'nexus-alarmas' (prioridad máxima con bypassDnd activado).",
            "Triggers Temporales: Notificaciones programadas a futuro tipo SchedulableTriggerInputTypes.DATE.",
            "Sincronización de IDs: El ID de la notificación se almacena en Firestore para permitir cancelaciones y reprogramaciones.",
            "Compatibilidad: Aislamiento dinámico para operar limpiamente en Expo Go y Development Builds."
        ],
        "speech": "Las notificaciones del sistema están centralizadas en services/notifications.js. Se crearon dos canales dedicados para Android: uno de prioridad regular para recordatorios y otro de prioridad máxima para alarmas, el cual tiene habilitado el traspaso del modo No Molestar (bypassDnd).\n\nCuando el usuario agenda una tarea o alarma, scheduleReminderNotification calcula la fecha exacta y la envía al planificador del sistema operativo con scheduleNotificationAsync. El ID generado por la notificación se almacena junto al documento en Firestore, lo que nos permite cancelar o reprogramar la notificación en el sistema operativo si el usuario borra la alarma o la marca como completada.",
        "codigo": [
            ("Definición y canales de notificación", "services/notifications.js (Líneas 22-46)"),
            ("Programación y cancelación horaria", "services/notifications.js (Líneas 74-115)"),
            ("Reconciliación y sincronización con Firestore", "context/RemindersContext.js (Líneas 100-127)"),
            ("Inicialización al montar la app", "App.js (Líneas 128-135)")
        ]
    },
    {
        "num": "4",
        "tema": "Respuesta por Voz de la IA (Text-to-Speech)",
        "slides": [
            "Motor Text-to-Speech: expo-speech.",
            "Parámetros de voz: Idioma español (es-ES), tono (pitch: 1.0), velocidad (rate: 1.0).",
            "Prevención de colisiones: Speech.stop() antes de cada emisión para evitar solapamientos.",
            "Integración en Chat: La función responderNexus() muestra el texto en pantalla y emite el audio simultáneamente.",
            "Configurabilidad: Interruptor interactivo para silenciar o activar la voz según la preferencia del usuario."
        ],
        "speech": "Para que el asistente NEXUS pudiera responder verbalmente, integramos la síntesis de voz en services/voice.js utilizando expo-speech.\n\nDiseñamos la función hablar(texto), la cual detiene cualquier reproducción anterior en cola para evitar que las frases se superpongan y sintetiza la respuesta con acento en español y modulación natural.\n\nEn ChatScreen.js, cada vez que la lógica del asistente genera una respuesta de confirmación, saludo o error, se ejecuta responderNexus(texto), que simultáneamente publica la burbuja en el chat y emite el audio si el usuario tiene el interruptor de voz habilitado.",
        "codigo": [
            ("Módulo de síntesis de voz (hablar / detener)", "services/voice.js (Líneas 30-51)"),
            ("Emisión de audio ligada al chat", "screens/ChatScreen.js (Líneas 66-69)")
        ]
    },
    {
        "num": "5",
        "tema": "Reconocimiento y Procesamiento de Órdenes por Voz",
        "slides": [
            "Speech-to-Text: expo-speech-recognition para captura de audio con soporte interimResults y eventos finales.",
            "Motor de PLN propio: services/assistant.js (sin dependencias externas pesadas).",
            "Detección de Intenciones: crear, eliminar, listar, hora, ayuda, saludo.",
            "Extracción temporal inteligente: Soporte para 'en 20 minutos', 'mañana a las 3 pm', 'el próximo viernes'.",
            "Limpieza de lenguaje natural: Filtra muletillas como 'oye nexus', 'por favor', 'recuérdame'."
        ],
        "speech": "La captura de órdenes por voz combina dos capas fundamentales:\n\n1. Reconocimiento del habla: En services/voice.js, mediante expo-speech-recognition, abrimos el micrófono del dispositivo, solicitamos los permisos requeridos y escuchamos en streaming. Cuando el usuario termina de hablar (isFinal), el texto capturado se envía al procesador.\n\n2. Procesamiento de Lenguaje Natural: En services/assistant.js desarrollamos el motor lógico procesarOrden(). Este analiza el texto, normaliza las palabras y extrae tanto la intención como las entidades de tiempo gracias a resolverFecha(). Por ejemplo, si el usuario dice: 'NEXUS, por favor recuérdame tomar agua en 15 minutos', el sistema extrae automáticamente el título 'Tomar agua', calcula la fecha sumando los minutos a la hora actual y ejecuta la creación del recordatorio en Firestore y la alarma en el sistema operativo.",
        "codigo": [
            ("Captura de micrófono y streaming", "services/voice.js (Líneas 64-152)"),
            ("Botón de micrófono e interacción de UI", "screens/ChatScreen.js (Líneas 143-178)"),
            ("Extracción de horas y fechas relativas", "services/assistant.js (Líneas 93-226)"),
            ("Motor de intenciones (procesarOrden)", "services/assistant.js (Líneas 305-374)")
        ]
    }
]

for sec in secciones:
    add_header_callout(f"TEMA {sec['num']}: {sec['tema']}")
    
    # Diapositiva
    dp = doc.add_paragraph()
    r = dp.add_run("📌 Qué colocar en la Diapositiva:")
    r.bold = True
    r.font.size = Pt(11)
    for b in sec['slides']:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(b)
        
    # Que decir
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(6)
    r2 = sp.add_run("🎙️ Qué decir durante la Exposición:")
    r2.bold = True
    r2.font.size = Pt(11)
    
    table_speech = doc.add_table(rows=1, cols=1)
    table_speech.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_s = table_speech.rows[0].cells[0]
    set_cell_background(cell_s, "F5F3FF")
    p_box = cell_s.paragraphs[0]
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(4)
    r_speech = p_box.add_run(sec['speech'])
    r_speech.font.size = Pt(10)
    r_speech.italic = True
    
    # Ubicacion en codigo
    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(8)
    r3 = cp.add_run("📂 Ubicación exacta en el Código:")
    r3.bold = True
    r3.font.size = Pt(11)
    
    table_code = doc.add_table(rows=len(sec['codigo']) + 1, cols=2)
    table_code.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_code.autofit = False
    
    hdr_cells = table_code.rows[0].cells
    hdr_cells[0].text = "Función / Componente"
    hdr_cells[1].text = "Archivo y Líneas"
    set_cell_background(hdr_cells[0], "EDE9FE")
    set_cell_background(hdr_cells[1], "EDE9FE")
    for cell in hdr_cells:
        for p in cell.paragraphs:
            for r_c in p.runs:
                r_c.bold = True
                r_c.font.size = Pt(10)
                
    for i, (comp, ubic) in enumerate(sec['codigo']):
        row_cells = table_code.rows[i+1].cells
        row_cells[0].text = comp
        row_cells[1].text = ubic
        set_cell_background(row_cells[0], "FAFAFA" if i%2==0 else "FFFFFF")
        set_cell_background(row_cells[1], "FAFAFA" if i%2==0 else "FFFFFF")
        for cell in row_cells:
            for p in cell.paragraphs:
                for r_c in p.runs:
                    r_c.font.size = Pt(9.5)
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

# Tabla resumen
add_header_callout("TABLA RESUMEN DE ARQUITECTURA")
t_res = doc.add_table(rows=6, cols=3)
t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Módulo", "Archivos Clave", "Tecnologías"]
for j, h in enumerate(headers):
    t_res.rows[0].cells[j].text = h
    set_cell_background(t_res.rows[0].cells[j], "6B21A8")
    for p in t_res.rows[0].cells[j].paragraphs:
        for r_h in p.runs:
            r_h.bold = True
            r_h.font.color.rgb = RGBColor(255, 255, 255)

filas = [
    ("Autenticación", "firebaseConfig.js, AuthContext.js, LoginScreen.js, RegisterScreen.js", "Firebase Auth, AsyncStorage"),
    ("Base de Datos", "firebaseConfig.js, RemindersContext.js, RegisterScreen.js, HomeScreen.js", "Cloud Firestore (onSnapshot, addDoc)"),
    ("Notificaciones", "services/notifications.js, RemindersContext.js, App.js", "expo-notifications, Android Notification Channels"),
    ("Voz IA (Salida)", "services/voice.js, ChatScreen.js", "expo-speech (Text-to-Speech es-ES)"),
    ("Órdenes de Voz (Entrada)", "services/voice.js, services/assistant.js, ChatScreen.js", "expo-speech-recognition, Regex PLN Engine")
]

for idx, (m, a, t) in enumerate(filas):
    row_c = t_res.rows[idx+1].cells
    row_c[0].text = m
    row_c[1].text = a
    row_c[2].text = t
    for cell in row_c:
        for p in cell.paragraphs:
            for r_c in p.runs:
                r_c.font.size = Pt(9.5)

output_path = r"c:\Users\ACER\Documents\programacionweb\NEXUS\Documentacion_Exposicion_NEXUS.docx"
doc.save(output_path)
print(f"Documento guardado exitosamente en: {output_path}")
